#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ai-video-sannong skill - API client(通用题材版)

单网关单 key 架构:
  - 全部走火山引擎 ARK(https://ark.cn-beijing.volces.com/api/v3)
  - 图片生成:Doubao-Seedream-5.0-lite(同步 /images/generations)
  - 视频生成:Doubao-Seedance-1.5-pro(异步任务 /contents/generations/tasks,原生配音)

支持的子命令:
  check-config              检查 ARK key / 用户偏好 / 预检记录
  save-key <KEY>            保存 ARK key(以 ark- 开头),会清空旧的预检记录
  save-prefs --json '...'   保存/合并用户创作偏好(方向、人群、音色、模式等)
  check-video-channel       预检视频模型可达性(支持 --model,结果记入 config)
  check-image-channel       预检图片模型可达性(支持 --model / --real)
  diagnose                  综合诊断:账号 + 网关可达性
  estimate-cost             预估成本(感知 --image-model / --video-model 单价)
  init-project              初始化项目目录 output/<日期-slug>/ + project.json
  gen-base-image            文生图(Seedream,默认 1080x1920 真 9:16,尺寸被拒自动回退)
  gen-variant-image         图生图编辑(Seedream,基于基准图)
  gen-video-clip            单镜图生视频(Seedance,异步轮询,默认带配音)
  gen-video-batch           多镜并发图生视频(读 jobs JSON,进度打 stderr)
  extract-frames            ffmpeg 抽帧(firstlast=承接自检 / fps=对标视频拆解)
  concat-clips              ffmpeg 把多段 mp4 无缝拼接成成片
  build-storyboard          生成最终交付 storyboard.xlsx

模型可通过 --model 参数覆盖,后续可平滑切换到 Seedance 2.0 / Seedream 标准版等。

调用约定:
  - 所有命令都返回 JSON 到 stdout(给 agent 解析)
  - 进度和警告走 stderr;错误信息走 stderr,退出码非 0
  - 配置文件路径:~/.ai-video-sannong/config.json (0600 权限)
  - Config 结构: {"ark_api_key": "ark-...", "preferences": {...}, "preflight": {...}}

依赖:
  pip install requests openpyxl
  ffmpeg(extract-frames / concat-clips 需要,系统命令)
"""

import argparse
import base64
import json
import os
import re
import sys
import threading
import time
from pathlib import Path

try:
    import requests
except ImportError:
    sys.stderr.write("ERROR: requests 包未安装。执行: pip install requests\n")
    sys.exit(2)


# ============================================================
# 常量
# ============================================================

ARK_BASE_URL = "https://ark.cn-beijing.volces.com/api/v3"

CONFIG_DIR = Path.home() / ".ai-video-sannong"
CONFIG_PATH = CONFIG_DIR / "config.json"

# 默认模型(都跑在火山 ARK 上,可通过 --model 覆盖)
MODEL_IMAGE = "doubao-seedream-5-0-lite"                   # 图片
MODEL_VIDEO_I2V = "doubao-seedance-1-5-pro-251215"         # 视频(原生支持音频)

MODEL_DISPLAY = {
    "doubao-seedream-5-0-lite": "Doubao-Seedream-5.0-Lite",
    "doubao-seedream-4-5-251128": "Doubao-Seedream-4.5",
    "doubao-seedream-4-0-250828": "Doubao-Seedream-4.0",
    "doubao-seedance-1-5-pro-251215": "Doubao-Seedance-1.5-Pro",
    "doubao-seedance-1-0-pro-250528": "Doubao-Seedance-1.0-Pro",
    "doubao-seedance-1-0-pro-fast-251015": "Doubao-Seedance-1.0-Pro-Fast",
}

# 成本估算(美元,以官方公开报价 + 实测估算,实际以账单为准)
IMAGE_MODEL_PRICING = {                                     # $/张
    "doubao-seedream-5-0-lite": 0.02,
    "doubao-seedream-4-5-251128": 0.03,
    "doubao-seedream-4-0-250828": 0.03,
}
VIDEO_MODEL_PRICING_PER_SEC = {                             # $/秒
    "doubao-seedance-1-5-pro-251215": 0.014,                # 含音频
    "doubao-seedance-1-0-pro-250528": 0.010,                # 无音频
    "doubao-seedance-1-0-pro-fast-251015": 0.008,           # 无音频
}
COST_IMAGE_DEFAULT = 0.02
COST_VIDEO_PER_SEC_DEFAULT = 0.014
CNY_PER_USD = 7.3

# 1.0 系列不支持音频生成
VIDEO_MODELS_WITHOUT_AUDIO_PREFIXES = ("doubao-seedance-1-0",)

# 图片尺寸:1080x1920 是真 9:16(注意 1024x1536 是 2:3,会被 Seedance 按 9:16
# 裁切或拉伸导致构图损失)。如果当前模型不支持该像素值,自动回退到比例式 "9:16"。
DEFAULT_IMAGE_SIZE = "1080x1920"
IMAGE_SIZE_FALLBACK = "9:16"

# 预检用的占位图尺寸(Seedance 对输入图有最小边长/宽高比要求,1×1 会被参数校验拒掉)
PROBE_IMAGE_W, PROBE_IMAGE_H = 432, 768                     # 9:16,最小边 432 ≥ 300

# 超时与轮询
HTTP_CONNECT_TIMEOUT = 30
HTTP_READ_TIMEOUT = 60
DOWNLOAD_TIMEOUT = 300
POLL_INTERVAL_SEC = 10
POLL_TIMEOUT_SEC = 600

DURATION_MIN, DURATION_MAX = 3, 15                          # 单镜秒数允许区间


# ============================================================
# 音色风格预设(Voice Presets)
# ============================================================
#
# Seedance 1.5 pro 不是传统 TTS,没有"音色 ID"参数。它是端到端音视频联合
# 生成模型,音色完全由 prompt 自然语言"召唤"。描述越骨架化,模型越退到
# 默认"标准播音腔"(僵硬);描述越具体(音色质感 + 情绪 + 语速 + 场景化
# 定位),模型生成的人声越自然贴合内容。
#
# preset 自带 gender 字段;显式传 --voice-gender 时以显式值为准。
# 方言不在 preset 里:用 --voice-style "亲切的中年女声,四川话" 之类自由描述。
# 详细选择逻辑见 references/voice-presets.md。

VOICE_PRESETS = {
    # 默认 preset:生活妙招类内容最通用
    "friendly-girl": {
        "label": "亲切邻家女孩(生活妙招通用)",
        "gender": "female",
        "style": "年轻、清亮、亲切、带一点轻松笑意的声线",
        "emotion": "轻松友好,像跟朋友分享小窍门",
        "pace": "中速自然,不刻意拖长",
    },
    # 适合传统手艺 / 古法配方 / 长辈代代相传感
    "warm-mama": {
        "label": "温暖妈妈(传统感强)",
        "gender": "female",
        "style": "温暖、成熟、有质感、稍低沉的中年声线",
        "emotion": "温暖耐心,像妈妈在教做饭的语气",
        "pace": "稍慢,字句之间留出呼吸感",
    },
    # 适合年轻向短视频 / 反差感强 / 种草
    "chirpy-host": {
        "label": "活泼达人(年轻向)",
        "gender": "female",
        "style": "年轻、活泼、精神感强、带轻微上扬尾音的声线",
        "emotion": "兴奋上扬,像短视频博主激情分享干货的语气",
        "pace": "略快,节奏感强",
    },
    # 适合科普向 / 原理讲解 / 知识带货
    "wisdom-aunt": {
        "label": "知性阿姨(科普感强)",
        "gender": "female",
        "style": "知性、沉稳、中年偏上、有质感不冰冷的声线",
        "emotion": "平静自信,像在传授知识又不端着的语气",
        "pace": "适中,信息密度高时稍慢一些保证清晰",
    },
    # 适合乡土 / 户外 / 硬核实操类内容
    "folksy-uncle": {
        "label": "接地气大叔(乡土/实操感强)",
        "gender": "male",
        "style": "厚实、接地气、嗓音略带沙哑的中年男声线",
        "emotion": "实在热心,像老乡在田埂上给你支招",
        "pace": "中速偏慢,句尾干脆利落",
    },
    # 适合科普测评 / 对比讲解 / 纪录片感内容
    "calm-narrator": {
        "label": "沉稳旁白(测评/科普男声)",
        "gender": "male",
        "style": "低沉、平稳、干净的男声线,带纪录片旁白质感",
        "emotion": "克制冷静,有分寸不夸张",
        "pace": "中速,重点词稍作停顿",
    },
}

VOICE_PRESET_DEFAULT = "friendly-girl"

# 这些词出现在音色描述里时,不再追加"女声/男声"后缀,避免"……男声线男声口播"式重复
GENDER_WORDS = ["女声", "男声", "女孩", "男孩", "女生", "男生",
                "大叔", "阿姨", "大妈", "阿伯", "男声线", "女声线"]


# ============================================================
# 输出工具
# ============================================================

_stderr_lock = threading.Lock()


def log(msg: str) -> None:
    """线程安全的 stderr 进度输出。"""
    with _stderr_lock:
        sys.stderr.write(msg.rstrip("\n") + "\n")
        sys.stderr.flush()


def emit(payload: dict) -> None:
    """把结构化结果输出到 stdout,供 agent 解析。"""
    sys.stdout.write(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    sys.stdout.flush()


def die(msg: str, code: int = 1, **extra) -> None:
    """打印错误到 stderr,带可选结构化字段,然后退出。"""
    payload = {"ok": False, "error": msg}
    payload.update(extra)
    sys.stderr.write(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    sys.exit(code)


class ClipError(Exception):
    """视频任务级错误。并发 batch 里逐 job 捕获,单镜命令里转成 die()。"""

    def __init__(self, msg: str, code: int = 1, **extra):
        super().__init__(msg)
        self.msg = msg
        self.code = code
        self.extra = extra


# ============================================================
# 配置管理
# ============================================================

def load_config() -> dict:
    if not CONFIG_PATH.exists():
        return {}
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        die(f"配置文件损坏: {e}", code=3)


def save_config(cfg: dict) -> None:
    CONFIG_DIR.mkdir(mode=0o700, exist_ok=True)
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    try:
        os.chmod(CONFIG_PATH, 0o600)
    except Exception:
        pass  # Windows 无 chmod 概念


def get_ark_key() -> str:
    """火山引擎 ARK key,用于图片和视频生成。"""
    cfg = load_config()
    key = cfg.get("ark_api_key", "").strip()
    if not key:
        die("ARK key 缺失,先运行 save-key ark-...", code=4)
    return key


def ark_headers(content_type: str = "application/json") -> dict:
    h = {"Authorization": f"Bearer {get_ark_key()}"}
    if content_type:
        h["Content-Type"] = content_type
    return h


def _record_preflight(kind: str, model: str, status: str, note: str = "") -> None:
    """把预检结果记入 config,供 check-config / Stage 3.5 跳过逻辑使用。"""
    cfg = load_config()
    pf = cfg.setdefault("preflight", {})
    pf[f"{kind}:{model}"] = {
        "status": status,
        "note": note,
        "at": time.strftime("%Y-%m-%d %H:%M:%S"),
    }
    save_config(cfg)


def _display(model: str) -> str:
    return MODEL_DISPLAY.get(model, model)


# ============================================================
# 错误分类(预检 + 生成共用)
# ============================================================

def _classify_ark_error(status_code: int, body: str) -> str:
    """把 ARK 返回的错误归类。

    返回值:auth_failed / quota / model_not_open / param_rejected / other
    注意:param_rejected(参数校验报错)说明请求已经过了鉴权和模型路由层,
    对"渠道预检"而言应视为渠道可用的信号,而不是渠道故障。
    """
    bl = body.lower()
    if status_code in (401, 403) or "unauthorized" in bl or "authentication" in bl:
        return "auth_failed"
    if "balance" in bl or "quota" in bl or "余额" in body or "insufficient" in bl:
        return "quota"
    if (("model" in bl and ("not" in bl or "disabled" in bl)) or "未开通" in body
            or status_code == 404):
        return "model_not_open"
    compact = bl.replace(" ", "").replace("_", "")
    if ("invalidparameter" in compact or "参数" in body
            or (("size" in bl or "image" in bl or "ratio" in bl or "duration" in bl
                 or "resolution" in bl)
                and ("invalid" in bl or "must" in bl or "should" in bl
                     or "expect" in bl or "support" in bl))):
        return "param_rejected"
    return "other"


def _diagnose_ark_error(status_code: int, body: str, model: str = "") -> str:
    """给 ARK 调用失败的常见错误一个友好的中文诊断提示。"""
    reason = _classify_ark_error(status_code, body)
    model_name = _display(model) if model else "对应模型"
    if reason == "auth_failed":
        return "【鉴权失败】ARK key 无效或已过期。重新运行 save-key ark-... 保存正确的 key。"
    if reason == "quota":
        return "【余额不足】账户额度不够。登录 https://console.volcengine.com/ark 充值。"
    if reason == "model_not_open":
        return (f"【模型未开通或模型名错误】检查 --model 拼写;若拼写无误,"
                f"登录 https://console.volcengine.com/ark/region:ark+cn-beijing/model"
                f",在「模型管理」里找到 {model_name} 并点击「开通」(免费)。")
    if reason == "param_rejected":
        return "【参数被拒】鉴权和模型都正常,是请求参数问题(尺寸/格式/时长等),看原始错误调整参数。"
    return ""


# ============================================================
# 占位图生成(预检用,纯标准库)
# ============================================================

def _make_solid_png(width: int, height: int, rgb=(180, 180, 180)) -> bytes:
    """生成一张纯色 PNG(不依赖 PIL),用于视频渠道预检的合规占位图。"""
    import struct
    import zlib

    def chunk(typ: bytes, data: bytes) -> bytes:
        c = struct.pack(">I", len(data)) + typ + data
        c += struct.pack(">I", zlib.crc32(typ + data) & 0xFFFFFFFF)
        return c

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)  # 8bit RGB
    row = b"\x00" + bytes(rgb) * width                            # filter 0 + 像素
    idat = zlib.compress(row * height, 9)
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", idat)
            + chunk(b"IEND", b""))


# ============================================================
# 子命令:check-config
# ============================================================

def cmd_check_config(_args) -> None:
    cfg = load_config()
    ark_key = cfg.get("ark_api_key", "").strip()
    preferences = cfg.get("preferences", {})
    preflight = cfg.get("preflight", {})

    if not ark_key:
        emit({
            "ok": True,
            "status": "missing",
            "missing": ["ark_api_key"],
            "preferences": preferences,
            "preflight": preflight,
            "message": "火山 ARK key 还没配,请运行 save-key ark-...",
        })
        return

    status = _quick_validate_ark(ark_key)
    emit({
        "ok": True,
        "status": "ready" if status["ok"] else "invalid",
        "message": ("火山 ARK key 已配置且可达" if status["ok"]
                    else "ARK key 验证失败,需要重新配置"),
        "ark_status": status,
        "preferences": preferences,
        "preflight": preflight,
    })


def _quick_validate_ark(key: str) -> dict:
    """火山 ARK 鉴权探测——用 GET /contents/generations/tasks?page_size=1 最轻量验证。"""
    try:
        r = requests.get(
            f"{ARK_BASE_URL}/contents/generations/tasks",
            headers={"Authorization": f"Bearer {key}"},
            params={"page_size": 1},
            timeout=15,
        )
        if r.status_code in (401, 403):
            return {"ok": False, "http_status": r.status_code,
                    "error": "ARK key 鉴权失败"}
        return {"ok": True, "http_status": r.status_code,
                "note": "鉴权通过(实际任务调用未测试)"}
    except requests.RequestException as e:
        return {"ok": False, "error": str(e)}


# ============================================================
# 子命令:save-key
# ============================================================

def cmd_save_key(args) -> None:
    """保存火山 ARK key。仅接受以 ark- 开头的 key。会清空旧 key 的预检记录。"""
    key = args.key.strip()
    if not key.startswith("ark-"):
        die(
            "ARK key 必须以 ark- 开头。从火山方舟控制台 "
            "https://console.volcengine.com/ark API Key 管理里复制。",
            code=6,
        )

    status = _quick_validate_ark(key)
    if not status["ok"]:
        die(f"ARK key 验证失败: {status.get('error', '未知错误')}", code=7,
            http_status=status.get("http_status"))

    cfg = load_config()
    cfg["ark_api_key"] = key
    cfg["saved_at"] = int(time.time())
    cfg.pop("preflight", None)  # 换 key 后预检记录作废
    save_config(cfg)

    emit({
        "ok": True,
        "message": "火山 ARK key 已验证并保存。建议紧接着跑 check-video-channel 和 "
                   "check-image-channel,把模型开通问题在配置阶段就暴露出来。",
        "config_path": str(CONFIG_PATH),
    })


# ============================================================
# 子命令:save-prefs(用户创作偏好记忆)
# ============================================================

def cmd_save_prefs(args) -> None:
    """保存/合并用户创作偏好。

    偏好是跨项目复用的"用户画像",建议字段(都可选):
      direction            题材方向,如 "三农园艺" / "家居清洁收纳"
      audience             目标人群,如 "都市阳台种菜白领"
      monetization         变现方向,如 "液态肥/营养液"
      voice_preset         默认音色 preset 名
      aspect_ratio         默认画幅,如 "9:16"
      mode                 默认确认模式 standard / fast / auto
      auto_approve_under_cny  低于该金额(人民币)的成本确认自动通过

    传 null 值可删除对应字段;--clear 清空全部偏好。
    """
    cfg = load_config()
    if args.clear:
        cfg.pop("preferences", None)
        save_config(cfg)
        emit({"ok": True, "message": "偏好已清空", "preferences": {}})
        return

    if not args.json:
        die("必须提供 --json '{...}' 或 --clear", code=6)
    try:
        data = json.loads(args.json)
    except Exception as e:
        die(f"--json 不是合法 JSON: {e}", code=6)
    if not isinstance(data, dict):
        die("--json 必须是一个 JSON 对象", code=6)

    prefs = cfg.get("preferences", {})
    for k, v in data.items():
        if v is None:
            prefs.pop(k, None)
        else:
            prefs[k] = v
    cfg["preferences"] = prefs
    save_config(cfg)
    emit({"ok": True, "message": "偏好已保存", "preferences": prefs})


# ============================================================
# 子命令:check-video-channel(渠道预检)
# ============================================================

def cmd_check_video_channel(args) -> None:
    """预检火山 ARK 视频生成渠道。

    ARK 没有公开的"模型列表"接口,所以做一次最轻量的 task 创建探测:
    用一张 432x768(9:16、满足最小边长要求)的纯色占位图 + 极短 prompt 提交任务,
    拿到 task_id 即证明渠道可用,随后**尽力取消该任务**(取消失败也不轮询不下载,
    任务自然过期,基本不计费)。

    判定规则:参数类报错(param_rejected)说明鉴权和模型路由都已通过,
    同样视为渠道可用——这避免了占位图被参数校验拒掉时的误报。
    """
    key = get_ark_key()
    model = args.model
    headers = {
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
    }

    probe_png = _make_solid_png(PROBE_IMAGE_W, PROBE_IMAGE_H)
    data_url = "data:image/png;base64," + base64.b64encode(probe_png).decode("ascii")

    probe_payload = {
        "model": model,
        "content": [
            {"type": "text", "text": "测试画面静止 --duration 5 --ratio 9:16 --watermark false"},
            {"type": "image_url", "image_url": {"url": data_url}},
        ],
    }

    try:
        r = requests.post(
            f"{ARK_BASE_URL}/contents/generations/tasks",
            headers=headers,
            json=probe_payload,
            timeout=HTTP_READ_TIMEOUT,
        )
    except requests.RequestException as e:
        emit({
            "ok": True,
            "status": "unknown",
            "reason": "probe_network_error",
            "message": f"网络错误,无法判断: {e}。可以直接跑生成测试。",
        })
        return

    if r.status_code == 200:
        try:
            task_id = r.json().get("id")
        except Exception:
            task_id = None
        cancelled = False
        if task_id:
            try:
                cr = requests.delete(
                    f"{ARK_BASE_URL}/contents/generations/tasks/{task_id}",
                    headers={"Authorization": f"Bearer {key}"},
                    timeout=15,
                )
                cancelled = cr.status_code in (200, 202, 204)
            except requests.RequestException:
                pass
        _record_preflight("video", model, "available", "probe task created")
        emit({
            "ok": True,
            "status": "available",
            "reason": "probe_task_created",
            "message": "视频渠道可用,探测任务已创建" + ("并已取消。" if cancelled else "(取消失败,将自然过期,不必处理)。"),
            "probe_task_id": task_id,
            "probe_task_cancelled": cancelled,
            "model": model,
            "ark_base_url": ARK_BASE_URL,
        })
        return

    body = r.text[:500]
    reason = _classify_ark_error(r.status_code, body)

    # 参数被拒 = 鉴权 + 模型路由都已通过 → 渠道视为可用
    if reason == "param_rejected":
        _record_preflight("video", model, "available", "param validation reached")
        emit({
            "ok": True,
            "status": "available",
            "reason": "param_validation_reached",
            "message": "视频渠道可用(探测请求到达参数校验层,说明鉴权和模型开通都正常;"
                       "占位图参数被拒不影响真实生成)。",
            "http_status": r.status_code,
            "raw_error": body,
            "model": model,
        })
        return

    if reason == "auth_failed":
        remediation = [
            "ARK key 鉴权失败。登录 https://console.volcengine.com/ark 查看你的 key 是否正确",
            "重新跑 save-key ark-... 保存新 key",
        ]
    elif reason == "model_not_open":
        remediation = [
            f"模型 {model} 未在你的账号开通(或模型 ID 拼写错误)。",
            "登录火山方舟控制台 https://console.volcengine.com/ark/region:ark+cn-beijing/model",
            f"在「模型管理」里找到 {_display(model)},点击「开通」(免费)",
        ]
    elif reason == "quota":
        remediation = [
            "账户额度不足。登录 https://console.volcengine.com/ark 充值",
        ]
    else:
        remediation = [
            f"查看原始错误: {body}",
            "如果是模型 ID 错误,可换用 --model 指定其他视频模型(注意 1.0 系列无音频)",
        ]

    _record_preflight("video", model, "unavailable", reason)
    emit({
        "ok": True,
        "status": "unavailable",
        "reason": reason,
        "message": "视频渠道不可用",
        "http_status": r.status_code,
        "raw_error": body,
        "model_attempted": model,
        "remediation": remediation,
    })


# ============================================================
# 子命令:check-image-channel(图片渠道预检)
# ============================================================

def cmd_check_image_channel(args) -> None:
    """预检火山 ARK 图片生成渠道。

    默认零成本探测:故意提交一个不合法的 size("64x64"),如果返回参数类
    错误,说明鉴权和模型路由都已通过 → 渠道可用;如果返回鉴权/未开通/余额
    错误,则如实报告。

    --real 模式:用默认尺寸真实生成一张图(约 $0.02-0.03),100% 确证渠道
    可用,同时顺带验证默认尺寸是否被该模型支持。
    """
    get_ark_key()  # 没 key 直接报错
    model = args.model

    if args.real:
        payload = {
            "model": model,
            "prompt": "一张浅灰色纯色背景测试图,无任何物体",
            "size": DEFAULT_IMAGE_SIZE,
            "response_format": "url",
            "watermark": False,
        }
    else:
        payload = {
            "model": model,
            "prompt": "测试",
            "size": "64x64",  # 故意低于最小尺寸,预期参数报错
            "response_format": "url",
            "watermark": False,
        }

    try:
        r = requests.post(
            f"{ARK_BASE_URL}/images/generations",
            headers=ark_headers(),
            json=payload,
            timeout=120,
        )
    except requests.RequestException as e:
        emit({
            "ok": True,
            "status": "unknown",
            "reason": "probe_network_error",
            "message": f"网络错误,无法判断: {e}。可以直接跑生成测试。",
        })
        return

    if r.status_code == 200:
        _record_preflight("image", model, "available",
                          "real generation ok" if args.real else "probe returned 200")
        emit({
            "ok": True,
            "status": "available",
            "reason": "generation_succeeded",
            "message": ("图片渠道可用,真实生成成功(已产生约 $0.02-0.03 费用,"
                        f"且默认尺寸 {DEFAULT_IMAGE_SIZE} 已确认可用)。" if args.real
                        else "图片渠道可用(探测请求被接受)。"),
            "model": model,
        })
        return

    body = r.text[:500]
    reason = _classify_ark_error(r.status_code, body)

    if reason == "param_rejected":
        _record_preflight("image", model, "available", "param validation reached")
        note = ""
        if args.real:
            note = (f"注意:--real 模式下默认尺寸 {DEFAULT_IMAGE_SIZE} 被参数校验拒绝,"
                    f"生成时请改用比例式 size(如 9:16),生图命令默认会自动回退。")
        emit({
            "ok": True,
            "status": "available",
            "reason": "param_validation_reached",
            "message": "图片渠道可用(请求到达参数校验层,说明鉴权和模型开通都正常)。"
                       "如需 100% 确证可加 --real 做一次真实生成(约 $0.02)。",
            "size_warning": note or None,
            "http_status": r.status_code,
            "raw_error": body,
            "model": model,
        })
        return

    if reason == "auth_failed":
        remediation = [
            "ARK key 鉴权失败。重新跑 save-key ark-... 保存正确的 key",
        ]
    elif reason == "model_not_open":
        remediation = [
            f"模型 {model} 未开通(或 ID 拼写错误)。",
            "登录 https://console.volcengine.com/ark/region:ark+cn-beijing/model",
            f"在「模型管理」里找到 {_display(model)},点击「开通」(免费)",
        ]
    elif reason == "quota":
        remediation = ["账户额度不足。登录 https://console.volcengine.com/ark 充值"]
    else:
        remediation = [f"查看原始错误: {body}"]

    _record_preflight("image", model, "unavailable", reason)
    emit({
        "ok": True,
        "status": "unavailable",
        "reason": reason,
        "message": "图片渠道不可用",
        "http_status": r.status_code,
        "raw_error": body,
        "model_attempted": model,
        "remediation": remediation,
    })


# ============================================================
# 子命令:diagnose(综合诊断)
# ============================================================

def cmd_diagnose(_args) -> None:
    """综合诊断:火山 ARK 配置 + 网关可达性 + 预检记录。"""
    result = {"ok": True, "checks": {}}

    cfg = load_config()
    ark_key = cfg.get("ark_api_key", "").strip()

    result["checks"]["config_path"] = str(CONFIG_PATH)
    result["checks"]["ark_key_present"] = bool(ark_key)
    result["checks"]["preferences"] = cfg.get("preferences", {})
    result["checks"]["preflight"] = cfg.get("preflight", {})

    if ark_key:
        try:
            r = requests.get(
                f"{ARK_BASE_URL}/contents/generations/tasks",
                headers={"Authorization": f"Bearer {ark_key}"},
                params={"page_size": 1},
                timeout=15,
            )
            if r.status_code in (401, 403):
                result["checks"]["ark_reachable"] = False
                result["checks"]["ark_error"] = f"鉴权失败 HTTP {r.status_code}"
            else:
                result["checks"]["ark_reachable"] = True
                result["checks"]["ark_http_status"] = r.status_code
        except requests.RequestException as e:
            result["checks"]["ark_reachable"] = False
            result["checks"]["ark_error"] = str(e)
    else:
        result["checks"]["ark_reachable"] = None

    summary = []
    if not ark_key:
        summary.append("⚠️  火山 ARK key 未配置,请先 save-key ark-...")
    elif result["checks"]["ark_reachable"]:
        summary.append("✅ 火山 ARK 网关可达")
        summary.append(f"   默认图片模型: {MODEL_IMAGE}")
        summary.append(f"   默认视频模型: {MODEL_VIDEO_I2V}(原生支持配音)")
        summary.append("   模型是否开通,用 check-video-channel / check-image-channel 探测确认")
    else:
        summary.append(f"❌ 火山 ARK 网关不可达: {result['checks'].get('ark_error')}")

    result["summary"] = summary
    result["recommended_image_model"] = MODEL_IMAGE
    result["recommended_video_model"] = MODEL_VIDEO_I2V
    emit(result)


# ============================================================
# 子命令:estimate-cost
# ============================================================

def cmd_estimate_cost(args) -> None:
    shots = args.shots
    duration_per_shot = args.duration
    if not (DURATION_MIN <= duration_per_shot <= DURATION_MAX):
        die(f"--duration 必须在 {DURATION_MIN}-{DURATION_MAX} 秒之间", code=6)

    image_model = args.image_model
    video_model = args.video_model
    img_price = IMAGE_MODEL_PRICING.get(image_model)
    vid_price = VIDEO_MODEL_PRICING_PER_SEC.get(video_model)
    unknown_pricing = []
    if img_price is None:
        img_price = COST_IMAGE_DEFAULT
        unknown_pricing.append(image_model)
    if vid_price is None:
        vid_price = COST_VIDEO_PER_SEC_DEFAULT
        unknown_pricing.append(video_model)

    img_calls = shots
    vid_calls = shots
    cost_img = img_calls * img_price
    cost_vid = vid_calls * duration_per_shot * vid_price
    total = cost_img + cost_vid

    total_video_sec = duration_per_shot * shots
    plan_note = ""
    if total_video_sec > 20:
        plan_note = f"⚠️ 总时长 {total_video_sec}s 偏长,镜头多/镜头长都会放大一致性风险,确认这是有意为之"
    elif total_video_sec < 10:
        plan_note = f"⚠️ 总时长 {total_video_sec}s 偏短,完播率友好但信息量有限"
    else:
        plan_note = f"{shots} 镜 × {duration_per_shot}s = {total_video_sec}s,适配短视频平台热门时长窗口"

    audio_note = None
    if video_model.startswith(VIDEO_MODELS_WITHOUT_AUDIO_PREFIXES):
        audio_note = f"⚠️ {video_model} 不支持音频生成,口播配音需要后期在剪映里补"

    emit({
        "ok": True,
        "shots": shots,
        "duration_per_shot_sec": duration_per_shot,
        "total_video_sec": total_video_sec,
        "image_calls": img_calls,
        "video_calls": vid_calls,
        "cost_image_usd_estimated": round(cost_img, 4),
        "cost_video_usd_estimated": round(cost_vid, 4),
        "total_usd_estimated": round(total, 4),
        "total_cny_estimate": round(total * CNY_PER_USD, 2),
        "models": {"image": image_model, "video": video_model},
        "unit_prices": {"image_usd_per_call": img_price,
                        "video_usd_per_sec": vid_price},
        "unknown_pricing_models": unknown_pricing or None,
        "plan_note": plan_note,
        "audio_note": audio_note,
        "note": (
            "成本是估算值,且未含自检不合格时的自动重试(通常 +0~2 次生图调用)。"
            "真实单价请到 https://console.volcengine.com/ark 控制台核对账单。"
        ),
    })


# ============================================================
# 子命令:init-project(项目目录 + 状态文件)
# ============================================================

def cmd_init_project(args) -> None:
    """初始化项目目录,避免多条视频互相覆盖,并落地 project.json 状态文件。

    目录结构:
      <base-dir>/<YYYYMMDD>-<slug>/
        ├── project.json     # 状态:阶段 / 模式 / 方向 / 锚定 / 文件清单
        ├── images/  clips/  prompts/  frames/
    """
    slug = re.sub(r"[^a-z0-9\-]+", "-", args.slug.lower()).strip("-") or "project"
    base = Path(args.base_dir)
    name = f"{time.strftime('%Y%m%d')}-{slug}"
    project_dir = base / name
    n = 2
    while project_dir.exists():
        project_dir = base / f"{name}-{n}"
        n += 1

    subdirs = ["images", "clips", "prompts", "frames"]
    for sub in subdirs:
        (project_dir / sub).mkdir(parents=True, exist_ok=True)

    project = {
        "title": args.title or args.slug,
        "slug": slug,
        "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "mode": args.mode,                       # standard / fast / auto
        "direction": args.direction or None,     # 题材方向
        "profile": args.profile or None,         # 使用的题材包文件(如 profiles/sannong.md)
        "stage": "S2_anchors",                   # init 发生在选题确认后、双锚定前
        "topic": args.topic or None,
        "anchors": {"visual": None, "narrative": None},
        "shots_file": "shots.json",              # Stage 3 确认后写入
        "deliverables": {},
        "notes": [],
    }
    pj_path = project_dir / "project.json"
    pj_path.write_text(json.dumps(project, ensure_ascii=False, indent=2),
                       encoding="utf-8")

    emit({
        "ok": True,
        "project_dir": str(project_dir),
        "project_json": str(pj_path),
        "subdirs": subdirs,
        "message": "项目目录已创建。后续所有产物写进该目录;阶段推进时由 agent 更新 project.json。",
    })


# ============================================================
# 图片生成核心(gen-base-image / gen-variant-image 共用)
# ============================================================

def _request_image(payload: dict, size_fallback: str, timeout: int) -> tuple:
    """POST /images/generations,带 size 不支持时的自动回退。

    返回 (image_bytes, source_url, size_used)。失败直接 die()。
    """
    sizes_tried = []
    while True:
        sizes_tried.append(payload.get("size"))
        try:
            r = requests.post(
                f"{ARK_BASE_URL}/images/generations",
                headers=ark_headers(),
                json=payload,
                timeout=timeout,
            )
        except requests.RequestException as e:
            die(f"网络错误: {e}", code=5)

        if r.status_code == 200:
            try:
                data = r.json()
            except Exception as e:
                die(f"响应不是 JSON: {e}", code=11, body=r.text[:300])
            image_bytes, image_url = _extract_image_bytes(data)
            return image_bytes, image_url, payload.get("size")

        body = r.text[:500]
        bl = body.lower()
        size_related = "size" in bl or "尺寸" in body
        if (size_related and size_fallback
                and size_fallback not in sizes_tried):
            log(f"[warn] 尺寸 {payload.get('size')} 被模型拒绝({body[:120]}),"
                f"自动回退到 {size_fallback} 重试")
            payload["size"] = size_fallback
            continue

        hint = _diagnose_ark_error(r.status_code, body, payload.get("model", ""))
        die(f"生图失败 HTTP {r.status_code}: {body} {hint}", code=8,
            http_status=r.status_code)


def cmd_gen_base_image(args) -> None:
    """文生图 - 火山 Seedream(基准图,锚定视觉 + 故事起点状态)。"""
    prompt = _read_prompt(args.prompt_file, args.prompt)

    payload = {
        "model": args.model,
        "prompt": prompt,
        "size": args.size,
        "response_format": "url",
        "watermark": False,
    }
    image_bytes, image_url, size_used = _request_image(
        payload, args.size_fallback, timeout=120)

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_bytes(image_bytes)

    emit({
        "ok": True,
        "output_path": str(out_path),
        "size_requested": args.size,
        "size_used": size_used,
        "bytes": len(image_bytes),
        "model": args.model,
        "source_url": image_url,
    })


def cmd_gen_variant_image(args) -> None:
    """图生图编辑 - 火山 Seedream(基于基准图生成变体镜头,保证一致性)。

    同 /images/generations 端点,加 `image` 字段传基准图 base64 data URL。
    """
    base_path = Path(args.base_image)
    if not base_path.exists():
        die(f"基准图不存在: {base_path}", code=9)

    prompt = _read_prompt(args.edit_prompt_file, args.edit_prompt)

    image_b64 = base64.b64encode(base_path.read_bytes()).decode("ascii")
    mime = _guess_mime(base_path)

    payload = {
        "model": args.model,
        "prompt": prompt,
        "image": f"data:{mime};base64,{image_b64}",
        "size": args.size,
        "response_format": "url",
        "watermark": False,
        "sequential_image_generation": "disabled",  # 只要单图,不要组图
    }
    image_bytes, image_url, size_used = _request_image(
        payload, args.size_fallback, timeout=180)

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_bytes(image_bytes)

    emit({
        "ok": True,
        "output_path": str(out_path),
        "size_requested": args.size,
        "size_used": size_used,
        "bytes": len(image_bytes),
        "model": args.model,
        "source_url": image_url,
    })


# ============================================================
# 视频生成核心(gen-video-clip / gen-video-batch 共用,异常用 ClipError)
# ============================================================

def _job_motion_prompt(job: dict) -> str:
    pf = job.get("motion_prompt_file")
    if pf:
        p = Path(pf)
        if not p.exists():
            raise ClipError(f"motion prompt 文件不存在: {p}", code=9)
        return p.read_text(encoding="utf-8").strip()
    inline = (job.get("motion_prompt") or "").strip()
    if inline:
        return inline
    raise ClipError("必须提供 motion_prompt_file 或 motion_prompt", code=16)


def _resolve_voice(job: dict) -> dict:
    """解析音色:细粒度参数 > preset 字段;gender 显式传入 > preset 自带 > female。"""
    preset_name = job.get("voice_preset") or VOICE_PRESET_DEFAULT
    if preset_name not in VOICE_PRESETS:
        raise ClipError(
            f"未知的 voice preset: {preset_name}。可用值: {', '.join(VOICE_PRESETS.keys())}",
            code=6,
        )
    preset = VOICE_PRESETS[preset_name]
    style_override = (job.get("voice_style") or "").strip()
    return {
        "preset": preset_name,
        "gender": job.get("voice_gender") or preset.get("gender", "female"),
        "style": style_override or preset["style"],
        "emotion": (job.get("voice_emotion") or "").strip() or preset["emotion"],
        "pace": (job.get("voice_pace") or "").strip() or preset["pace"],
    }


def _build_video_prompt(job: dict, voice: dict) -> tuple:
    """拼接完整 prompt(三段:动作描述 + 音频指令 + Seedance CLI 参数)。

    返回 (full_prompt, audio_instruction)。
    """
    motion_prompt = _job_motion_prompt(job)
    prompt_parts = [motion_prompt.rstrip()]

    voiceover_text = (job.get("voiceover_text") or "").strip()
    generate_audio = job.get("generate_audio", True)
    audio_instruction = ""

    if generate_audio and voiceover_text:
        safe_text = voiceover_text.replace('"', "'").replace("\n", " ")
        override = (job.get("voice_prompt_override") or "").strip()
        if override:
            audio_instruction = override
        else:
            # 音色描述里已含性别字眼时不再追加"女声/男声"后缀,避免重复
            style_has_gender = any(k in voice["style"] for k in GENDER_WORDS)
            if style_has_gender:
                voice_intro = f"AI {voice['style']}口播"
            else:
                gender_zh = "女声" if voice["gender"] == "female" else "男声"
                voice_intro = f"AI {voice['style']}{gender_zh}口播"
            # 富信息描述(音色/情绪/语速/台词 4 维),骨架化指令会让模型退到僵硬播音腔
            audio_instruction = (
                f"音频:{voice_intro},"
                f"情绪{voice['emotion']},"
                f"语速{voice['pace']},"
                f"台词:'{safe_text}'"
            )
        prompt_parts.append(audio_instruction)

    # Seedance CLI 参数(模型协议关键字,保留英文)
    cli_params = (
        f"--resolution {job.get('resolution', '1080p')} "
        f"--ratio {job.get('aspect_ratio', '9:16')} "
        f"--duration {job.get('duration', 5)} "
        f"--camerafixed {str(job.get('camerafixed', True)).lower()} "
        f"--watermark {str(job.get('watermark', False)).lower()}"
    )
    prompt_parts.append(cli_params)
    return " ".join(prompt_parts), audio_instruction


def _run_video_job(job: dict) -> dict:
    """跑一个完整的单镜视频任务:校验 → 建任务 → 轮询 → 下载。

    job 字段(与 gen-video-clip 的 CLI 参数同名,下划线风格):
      image / motion_prompt_file / motion_prompt / output / duration /
      aspect_ratio / resolution / camerafixed / watermark / generate_audio /
      voiceover_text / voice_gender / voice_preset / voice_style /
      voice_emotion / voice_pace / voice_prompt_override / model / label
    """
    label = job.get("label") or Path(job.get("output", "clip")).stem

    image_path = Path(job.get("image", ""))
    if not image_path.exists():
        raise ClipError(f"分镜静图不存在: {image_path}", code=9)

    duration = int(job.get("duration", 5))
    if not (DURATION_MIN <= duration <= DURATION_MAX):
        raise ClipError(f"duration 必须在 {DURATION_MIN}-{DURATION_MAX} 秒之间,"
                        f"当前 {duration}", code=6)
    job["duration"] = duration

    model = job.get("model") or MODEL_VIDEO_I2V
    generate_audio = job.get("generate_audio", True)
    voiceover_text = (job.get("voiceover_text") or "").strip()

    if (generate_audio and voiceover_text
            and model.startswith(VIDEO_MODELS_WITHOUT_AUDIO_PREFIXES)):
        log(f"[{label}] ⚠️ 模型 {model} 不支持音频,voiceover 将被忽略(只有画面)")

    voice = _resolve_voice(job)
    full_prompt, audio_instruction = _build_video_prompt(job, voice)

    image_b64 = base64.b64encode(image_path.read_bytes()).decode("ascii")
    mime = _guess_mime(image_path)

    payload = {
        "model": model,
        "content": [
            {"type": "text", "text": full_prompt},
            {"type": "image_url",
             "image_url": {"url": f"data:{mime};base64,{image_b64}"}},
        ],
    }
    if generate_audio:
        payload["generate_audio"] = True

    headers = ark_headers()

    # === 步骤 1:创建任务 ===
    try:
        r = requests.post(
            f"{ARK_BASE_URL}/contents/generations/tasks",
            headers=headers,
            json=payload,
            timeout=(HTTP_CONNECT_TIMEOUT, HTTP_READ_TIMEOUT),
        )
    except requests.RequestException as e:
        raise ClipError(f"创建视频任务网络错误: {e}", code=5)

    if r.status_code != 200:
        body = r.text[:500]
        hint = _diagnose_ark_error(r.status_code, body, model)
        raise ClipError(f"创建视频任务失败 HTTP {r.status_code}: {body} {hint}",
                        code=10, http_status=r.status_code)

    try:
        task = r.json()
    except Exception as e:
        raise ClipError(f"任务响应不是 JSON: {e}", code=11, body=r.text[:300])

    task_id = task.get("id")
    if not task_id:
        raise ClipError(f"响应里没有 task id: {task}", code=11)

    log(f"[{label}] 任务已创建 id={task_id},开始轮询...")

    # === 步骤 2:轮询 ===
    start = time.time()
    last_status = None
    while True:
        elapsed = time.time() - start
        if elapsed > POLL_TIMEOUT_SEC:
            raise ClipError(f"轮询超时({POLL_TIMEOUT_SEC}s),任务 {task_id} 仍未完成",
                            code=12, task_id=task_id, last_status=last_status)
        time.sleep(POLL_INTERVAL_SEC)

        try:
            qr = requests.get(
                f"{ARK_BASE_URL}/contents/generations/tasks/{task_id}",
                headers={"Authorization": headers["Authorization"]},
                timeout=HTTP_READ_TIMEOUT,
            )
        except requests.RequestException as e:
            log(f"[{label}] [warn] 查询网络抖动: {e},继续轮询")
            continue

        if qr.status_code != 200:
            log(f"[{label}] [warn] 查询 HTTP {qr.status_code}: {qr.text[:200]}")
            continue

        try:
            info = qr.json()
        except Exception:
            continue

        status = info.get("status")
        if status != last_status:
            log(f"[{label}] status={status} ({int(elapsed)}s)")
            last_status = status

        # ARK Seedance 状态机:queued / running / succeeded / failed / cancelled
        if status == "succeeded":
            content = info.get("content") or {}
            video_url = content.get("video_url") or info.get("video_url")
            if not video_url:
                raise ClipError(f"任务成功但没找到 video_url。info: {info}", code=13)
            out_path = Path(job["output"])
            out_path.parent.mkdir(parents=True, exist_ok=True)
            try:
                vr = requests.get(video_url, timeout=DOWNLOAD_TIMEOUT, stream=True)
                vr.raise_for_status()
                with open(out_path, "wb") as f:
                    for chunk in vr.iter_content(chunk_size=65536):
                        if chunk:
                            f.write(chunk)
            except requests.RequestException as e:
                raise ClipError(f"下载视频失败: {e}", code=14, video_url=video_url)

            return {
                "ok": True,
                "label": label,
                "task_id": task_id,
                "output_path": str(out_path),
                "video_url": video_url,
                "duration_requested": duration,
                "elapsed_sec": round(elapsed, 1),
                "bytes": out_path.stat().st_size,
                "model": model,
                "audio": {
                    "generate_audio": generate_audio,
                    "has_voiceover": bool(voiceover_text),
                    "voiceover_text": voiceover_text or None,
                    "voice_gender": voice["gender"] if voiceover_text else None,
                    "voice_preset": voice["preset"] if voiceover_text else None,
                    "voice_style_final": voice["style"] if voiceover_text else None,
                    "voice_emotion_final": voice["emotion"] if voiceover_text else None,
                    "voice_pace_final": voice["pace"] if voiceover_text else None,
                    "audio_instruction": audio_instruction or None,
                },
                "usage": info.get("usage") or {},
            }

        if status == "failed":
            err = info.get("error") or info.get("failure_reason") or info
            raise ClipError(f"任务失败: {err}", code=15, task_id=task_id, error=str(err))

        if status == "cancelled":
            raise ClipError("任务被取消", code=16, task_id=task_id)

        # queued / running → 继续轮询


def _job_from_args(args) -> dict:
    return {
        "label": Path(args.output).stem,
        "image": args.image,
        "motion_prompt_file": args.motion_prompt_file,
        "motion_prompt": args.motion_prompt,
        "output": args.output,
        "duration": args.duration,
        "aspect_ratio": args.aspect_ratio,
        "resolution": args.resolution,
        "camerafixed": args.camerafixed,
        "watermark": args.watermark,
        "generate_audio": args.generate_audio,
        "voiceover_text": args.voiceover_text,
        "voice_gender": args.voice_gender,
        "voice_preset": args.voice_preset,
        "voice_style": args.voice_style,
        "voice_emotion": args.voice_emotion,
        "voice_pace": args.voice_pace,
        "voice_prompt_override": args.voice_prompt_override,
        "model": args.model,
    }


def cmd_gen_video_clip(args) -> None:
    """单镜图生视频(重做单镜 / 不想并发时用)。"""
    try:
        result = _run_video_job(_job_from_args(args))
    except ClipError as e:
        die(e.msg, code=e.code, **e.extra)
    emit(result)


def cmd_gen_video_batch(args) -> None:
    """多镜并发图生视频。

    读取 jobs JSON(数组,每项字段同 gen-video-clip 参数的下划线风格),
    batch 级 CLI 参数作为各 job 的默认值,job 内字段优先。
    各任务并发提交 + 轮询,进度逐行打到 stderr,最终汇总 JSON 到 stdout。

    任何 job 失败不影响其他 job;全部完成后若有失败,退出码 15。
    """
    jobs_path = Path(args.jobs_json)
    if not jobs_path.exists():
        die(f"jobs 文件不存在: {jobs_path}", code=9)
    try:
        jobs_raw = json.loads(jobs_path.read_text(encoding="utf-8"))
    except Exception as e:
        die(f"jobs 文件不是合法 JSON: {e}", code=6)
    if not isinstance(jobs_raw, list) or not jobs_raw:
        die("jobs 文件必须是非空 JSON 数组", code=6)

    defaults = {
        "duration": args.duration,
        "aspect_ratio": args.aspect_ratio,
        "resolution": args.resolution,
        "camerafixed": args.camerafixed,
        "watermark": args.watermark,
        "generate_audio": args.generate_audio,
        "voice_gender": args.voice_gender,
        "voice_preset": args.voice_preset,
        "model": args.model,
    }

    prepared = []
    for i, job in enumerate(jobs_raw, start=1):
        if not isinstance(job, dict):
            die(f"jobs[{i - 1}] 不是 JSON 对象", code=6)
        merged = dict(defaults)
        merged.update(job)
        merged.setdefault("label",
                          Path(merged.get("output", f"clip-{i:02d}")).stem)
        prepared.append(merged)

    log(f"[batch] 共 {len(prepared)} 个任务,max_workers={args.max_workers},开始并发提交...")

    from concurrent.futures import ThreadPoolExecutor, as_completed

    results = [None] * len(prepared)

    def run_one(idx: int, job: dict):
        try:
            return idx, _run_video_job(job)
        except ClipError as e:
            payload = {"ok": False, "label": job.get("label"),
                       "error": e.msg}
            payload.update(e.extra)
            return idx, payload
        except Exception as e:  # 防御性兜底,单 job 崩溃不拖垮整个 batch
            return idx, {"ok": False, "label": job.get("label"),
                         "error": f"未预期异常: {e}"}

    with ThreadPoolExecutor(max_workers=args.max_workers) as ex:
        futures = [ex.submit(run_one, i, j) for i, j in enumerate(prepared)]
        for fut in as_completed(futures):
            idx, res = fut.result()
            results[idx] = res
            mark = "✅ 完成" if res.get("ok") else f"❌ 失败: {res.get('error', '')[:120]}"
            log(f"[batch] {res.get('label')} {mark}")

    succeeded = sum(1 for r in results if r and r.get("ok"))
    emit({
        "ok": succeeded == len(results),
        "total": len(results),
        "succeeded": succeeded,
        "failed": len(results) - succeeded,
        "results": results,
    })
    if succeeded != len(results):
        sys.exit(15)


# ============================================================
# 子命令:extract-frames(ffmpeg 抽帧)
# ============================================================

def _require_ffmpeg() -> str:
    import shutil
    ffmpeg_bin = shutil.which("ffmpeg")
    if not ffmpeg_bin:
        die(
            "ffmpeg 不在系统 PATH 里。安装方式:\n"
            "  - macOS: brew install ffmpeg\n"
            "  - Ubuntu/Debian: sudo apt install ffmpeg\n"
            "  - Windows: 去 https://www.gyan.dev/ffmpeg/builds/ 下载,把 bin 加到 PATH",
            code=20,
        )
    return ffmpeg_bin


def cmd_extract_frames(args) -> None:
    """ffmpeg 抽帧。两种模式:

    firstlast:抽视频第一帧 + 最后一帧,用于多镜承接自检
              (镜 K 的尾帧 vs 镜 K+1 的首帧应该状态衔接)。
    fps:      按 --fps 频率抽帧(默认 1 帧/秒),用于对标视频拆解
              (agent 逐帧视觉分析,反推视觉锚定方案和镜头结构)。
    """
    import subprocess

    ffmpeg_bin = _require_ffmpeg()
    video_path = Path(args.video)
    if not video_path.exists():
        die(f"视频文件不存在: {video_path}", code=9)

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = video_path.stem
    produced = []

    def run_ffmpeg(cmd: list, timeout: int = 60) -> bool:
        try:
            res = subprocess.run(cmd, capture_output=True, text=True,
                                 timeout=timeout)
            return res.returncode == 0
        except subprocess.TimeoutExpired:
            return False

    if args.mode == "firstlast":
        first_path = out_dir / f"{stem}-first.png"
        ok_first = run_ffmpeg([
            ffmpeg_bin, "-y", "-i", str(video_path),
            "-frames:v", "1", str(first_path),
        ])
        if ok_first and first_path.exists() and first_path.stat().st_size > 0:
            produced.append(str(first_path))

        last_path = out_dir / f"{stem}-last.png"
        ok_last = False
        for sseof in ("-0.3", "-1.0"):  # 个别封装下 -0.3 取不到帧,回退 -1.0
            ok_last = run_ffmpeg([
                ffmpeg_bin, "-y", "-sseof", sseof, "-i", str(video_path),
                "-update", "1", "-frames:v", "1", str(last_path),
            ])
            if ok_last and last_path.exists() and last_path.stat().st_size > 0:
                produced.append(str(last_path))
                break
        if len(produced) < 2:
            die(f"抽帧失败(成功 {len(produced)}/2),检查视频文件是否完整",
                code=21, produced=produced)
    else:  # fps 模式
        pattern = out_dir / f"{stem}-%03d.png"
        ok = run_ffmpeg([
            ffmpeg_bin, "-y", "-i", str(video_path),
            "-vf", f"fps={args.fps}", str(pattern),
        ], timeout=180)
        produced = sorted(str(p) for p in out_dir.glob(f"{stem}-[0-9][0-9][0-9].png"))
        if not ok or not produced:
            die("按 fps 抽帧失败,检查视频文件是否完整", code=21)

    emit({
        "ok": True,
        "mode": args.mode,
        "video": str(video_path),
        "frames": produced,
        "count": len(produced),
        "note": ("用 Read/视觉能力逐帧查看这些图,firstlast 用于镜头承接自检,"
                 "fps 用于对标视频拆解。"),
    })


# ============================================================
# 子命令:concat-clips(ffmpeg 拼接成片)
# ============================================================

def cmd_concat_clips(args) -> None:
    """用 ffmpeg 把 N 段 mp4 拼接成一段成片。

    优先 concat demuxer + stream copy(不重新编码,无损且快);
    失败时 fallback 到 H.264 重新编码(慢但稳)。
    """
    import shutil
    import subprocess
    import tempfile

    ffmpeg_bin = _require_ffmpeg()

    clip_paths = [Path(p) for p in args.clips]
    missing = [str(p) for p in clip_paths if not p.exists()]
    if missing:
        die(f"以下视频文件不存在: {missing}", code=9)
    if len(clip_paths) < 2:
        die(f"至少需要 2 段视频才能拼接,当前只有 {len(clip_paths)} 段", code=9)

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    def write_concat_list() -> Path:
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".txt", delete=False, encoding="utf-8"
        ) as f:
            for clip in clip_paths:
                # concat demuxer 要求 file 'path' 格式,路径中的单引号要转义
                safe_path = str(clip.resolve()).replace("'", r"'\''")
                f.write(f"file '{safe_path}'\n")
            return Path(f.name)

    # 第一次尝试:stream copy
    list_file = write_concat_list()
    try:
        result = subprocess.run(
            [
                ffmpeg_bin, "-y",
                "-f", "concat", "-safe", "0",
                "-i", str(list_file),
                "-c", "copy",
                "-movflags", "+faststart",
                str(out_path),
            ],
            capture_output=True, text=True, timeout=120,
        )
    except subprocess.TimeoutExpired:
        die("ffmpeg 拼接超时(>120s),输入文件可能损坏或太大", code=21)
    finally:
        list_file.unlink(missing_ok=True)

    if result.returncode != 0:
        # stream copy 失败:多半是输入视频参数不一致,fallback 重新编码
        log(f"[warn] stream copy 拼接失败,尝试重新编码 fallback...\n"
            f"原始 ffmpeg 错误: {result.stderr[:500]}")
        list_file = write_concat_list()
        try:
            result = subprocess.run(
                [
                    ffmpeg_bin, "-y",
                    "-f", "concat", "-safe", "0",
                    "-i", str(list_file),
                    "-c:v", "libx264", "-preset", "fast", "-crf", "20",
                    "-c:a", "aac", "-b:a", "128k",
                    "-pix_fmt", "yuv420p",
                    "-movflags", "+faststart",
                    str(out_path),
                ],
                capture_output=True, text=True, timeout=300,
            )
        except subprocess.TimeoutExpired:
            die("ffmpeg 重新编码超时(>300s)", code=22)
        finally:
            list_file.unlink(missing_ok=True)
        if result.returncode != 0:
            die(
                f"ffmpeg 拼接失败(stream copy 和 reencode 都失败): {result.stderr[:500]}",
                code=22,
            )

    if not out_path.exists() or out_path.stat().st_size == 0:
        die("ffmpeg 报告成功但输出文件不存在或为空", code=23)

    # 用 ffprobe 探测输出时长(可选)
    output_duration_sec = None
    ffprobe_bin = shutil.which("ffprobe")
    if ffprobe_bin:
        try:
            probe = subprocess.run(
                [
                    ffprobe_bin, "-v", "error",
                    "-show_entries", "format=duration",
                    "-of", "default=noprint_wrappers=1:nokey=1",
                    str(out_path),
                ],
                capture_output=True, text=True, timeout=30,
            )
            if probe.returncode == 0 and probe.stdout.strip():
                output_duration_sec = round(float(probe.stdout.strip()), 2)
        except Exception:
            pass

    emit({
        "ok": True,
        "output_path": str(out_path),
        "input_clip_count": len(clip_paths),
        "bytes": out_path.stat().st_size,
        "output_duration_sec": output_duration_sec,
        "ffmpeg_bin": ffmpeg_bin,
        "note": (
            "成片已生成。可直接发布到抖音/小红书/视频号,"
            "也可以再丢到剪映里加 BGM/字幕。"
        ),
    })


# ============================================================
# 子命令:build-storyboard
# ============================================================

def cmd_build_storyboard(args) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        die("openpyxl 未安装。执行: pip install openpyxl", code=2)

    shots_json_path = Path(args.shots_json)
    if not shots_json_path.exists():
        die(f"分镜数据文件不存在: {shots_json_path}", code=9)
    shots = json.loads(shots_json_path.read_text(encoding="utf-8"))

    wb = Workbook()
    ws = wb.active
    ws.title = "分镜表"

    headers = ["镜头号", "时长(秒)", "前镜承接状态", "口播文案", "屏幕字幕",
               "视觉描述", "生图 Prompt", "生视频 Motion Prompt",
               "图片文件", "视频文件"]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E7D32")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)

    for i, shot in enumerate(shots, start=1):
        ws.append([
            i,
            shot.get("duration_sec", 5),
            shot.get("prev_state", "(起点,无前镜)" if i == 1 else ""),
            shot.get("voiceover", ""),
            shot.get("subtitle", ""),
            shot.get("visual_description_cn", ""),
            shot.get("image_prompt", ""),
            shot.get("motion_prompt", ""),
            shot.get("image_file", ""),
            shot.get("video_file", ""),
        ])

    widths = [8, 10, 22, 24, 18, 28, 40, 40, 18, 18]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = w

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 90

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border

    # Sheet 2: 口播稿
    ws2 = wb.create_sheet("口播稿")
    ws2["A1"] = "镜头号"
    ws2["B1"] = "口播文案"
    ws2["A1"].font = ws2["B1"].font = header_font
    ws2["A1"].fill = ws2["B1"].fill = header_fill
    for i, shot in enumerate(shots, start=1):
        ws2.cell(row=i + 1, column=1, value=i)
        ws2.cell(row=i + 1, column=2, value=shot.get("voiceover", ""))
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 40

    out_path = Path(args.output)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    # 同时输出纯文本口播稿
    vo_path = out_path.parent / "voiceover-script.txt"
    with open(vo_path, "w", encoding="utf-8") as f:
        for i, shot in enumerate(shots, start=1):
            f.write(f"【镜头 {i}】{shot.get('voiceover', '')}\n")

    emit({
        "ok": True,
        "storyboard_xlsx": str(out_path),
        "voiceover_txt": str(vo_path),
        "shots_count": len(shots),
    })


# ============================================================
# 内部工具
# ============================================================

def _read_prompt(prompt_file: str, inline_prompt: str) -> str:
    if prompt_file:
        return Path(prompt_file).read_text(encoding="utf-8").strip()
    if inline_prompt:
        return inline_prompt
    die("必须提供 --prompt-file 或 --prompt", code=16)


def _extract_image_bytes(resp: dict) -> tuple:
    """从 OpenAI 风格的 images 响应里取出图片字节。
    兼容 b64_json 和 url 两种返回。返回 (bytes, source_url_or_None)。"""
    if not isinstance(resp, dict):
        die(f"非法响应: {resp}", code=17)
    items = resp.get("data") or []
    if not items:
        die(f"响应里没有 data 字段: {resp}", code=17)
    first = items[0]
    if "b64_json" in first:
        return base64.b64decode(first["b64_json"]), None
    if "url" in first:
        url = first["url"]
        try:
            r = requests.get(url, timeout=DOWNLOAD_TIMEOUT)
            r.raise_for_status()
            return r.content, url
        except requests.RequestException as e:
            die(f"下载图片失败: {e}", code=18, url=url)
    die(f"data[0] 既没有 b64_json 也没有 url: {first}", code=17)


def _guess_mime(path: Path) -> str:
    suffix = path.suffix.lower()
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
    }.get(suffix, "image/png")


def _bool_arg(x) -> bool:
    return str(x).lower() == "true"


# ============================================================
# 入口
# ============================================================

def _add_video_job_args(sp, with_voiceover_text: bool = True) -> None:
    """gen-video-clip / gen-video-batch 共享的参数定义。"""
    sp.add_argument("--duration", type=int, default=5,
                    help=f"视频秒数(默认 5,允许 {DURATION_MIN}-{DURATION_MAX})")
    sp.add_argument("--aspect-ratio", default="9:16",
                    help="画幅比(默认 9:16 竖屏)")
    sp.add_argument("--resolution", default="1080p",
                    help="分辨率(默认 1080p)")
    sp.add_argument("--camerafixed", default=True, type=_bool_arg,
                    help="是否锁机位(默认 true)")
    sp.add_argument("--watermark", default=False, type=_bool_arg,
                    help="是否带水印(默认 false)")
    sp.add_argument("--generate-audio", default=True, type=_bool_arg,
                    help="是否生成音频(默认 true,仅 Seedance 1.5 pro 及以上支持)")
    sp.add_argument("--voice-gender", default=None,
                    choices=["female", "male"],
                    help="配音性别;不传时用 preset 自带性别(默认 preset 为女声)")
    sp.add_argument("--voice-preset", default=VOICE_PRESET_DEFAULT,
                    choices=list(VOICE_PRESETS.keys()),
                    help=(f"音色风格预设(默认 {VOICE_PRESET_DEFAULT})。可选: "
                          + ", ".join(f"{k}={v['label']}"
                                      for k, v in VOICE_PRESETS.items())))
    sp.add_argument("--model", default=MODEL_VIDEO_I2V,
                    help=f"视频模型 ID,默认 {MODEL_VIDEO_I2V}")
    if with_voiceover_text:
        sp.add_argument("--voiceover-text", default="",
                        help="当前镜的中文口播台词;非空时作为配音内容")
        sp.add_argument("--voice-style", default="",
                        help="覆盖 preset 的音色描述,如 '亲切的中年女声,四川话'(方言走这个)")
        sp.add_argument("--voice-emotion", default="",
                        help="覆盖 preset 的情绪基调,如 '神秘悬念'")
        sp.add_argument("--voice-pace", default="",
                        help="覆盖 preset 的语速,如 '缓慢一字一顿'")
        sp.add_argument("--voice-prompt-override", default="",
                        help="[逃生舱] 完全自定义整段音频指令,忽略 preset 和细粒度参数")


def main():
    p = argparse.ArgumentParser(description="ai-video-sannong API client(通用题材版)")
    sub = p.add_subparsers(dest="cmd", required=True)

    sp = sub.add_parser("check-config",
                        help="检查 ARK key / 偏好 / 预检记录")
    sp.set_defaults(func=cmd_check_config)

    sp = sub.add_parser("save-key",
                        help="保存火山 ARK key(必须以 ark- 开头)")
    sp.add_argument("key", help="火山 ARK API key,以 ark- 开头")
    sp.set_defaults(func=cmd_save_key)

    sp = sub.add_parser("save-prefs",
                        help="保存/合并用户创作偏好(JSON 对象)")
    sp.add_argument("--json", default="",
                    help='偏好 JSON,如 \'{"direction":"三农园艺","mode":"fast"}\';值为 null 删除该字段')
    sp.add_argument("--clear", action="store_true", help="清空全部偏好")
    sp.set_defaults(func=cmd_save_prefs)

    sp = sub.add_parser("check-video-channel",
                        help="预检视频模型可达性(结果记入 config)")
    sp.add_argument("--model", default=MODEL_VIDEO_I2V,
                    help=f"要探测的视频模型 ID(默认 {MODEL_VIDEO_I2V})")
    sp.set_defaults(func=cmd_check_video_channel)

    sp = sub.add_parser("check-image-channel",
                        help="预检图片模型可达性(默认零成本探测)")
    sp.add_argument("--model", default=MODEL_IMAGE,
                    help=f"要探测的图片模型 ID(默认 {MODEL_IMAGE})")
    sp.add_argument("--real", action="store_true",
                    help="真实生成一张测试图确证(约 $0.02-0.03,顺带验证默认尺寸)")
    sp.set_defaults(func=cmd_check_image_channel)

    sp = sub.add_parser("diagnose",
                        help="综合诊断:ARK key + 网关可达性 + 预检记录")
    sp.set_defaults(func=cmd_diagnose)

    sp = sub.add_parser("estimate-cost",
                        help="预估成本(感知模型单价)")
    sp.add_argument("--shots", type=int, default=3, help="镜头数(默认 3)")
    sp.add_argument("--duration", type=int, default=5,
                    help=f"单镜秒数(默认 5,允许 {DURATION_MIN}-{DURATION_MAX})")
    sp.add_argument("--image-model", default=MODEL_IMAGE,
                    help=f"图片模型(默认 {MODEL_IMAGE})")
    sp.add_argument("--video-model", default=MODEL_VIDEO_I2V,
                    help=f"视频模型(默认 {MODEL_VIDEO_I2V})")
    sp.set_defaults(func=cmd_estimate_cost)

    sp = sub.add_parser("init-project",
                        help="初始化项目目录 + project.json 状态文件")
    sp.add_argument("--slug", required=True,
                    help="项目短名,小写字母/数字/短横线(中文选题请转成拼音或英文)")
    sp.add_argument("--title", default="", help="项目标题(可中文)")
    sp.add_argument("--mode", default="standard",
                    choices=["standard", "fast", "auto"],
                    help="确认模式(默认 standard)")
    sp.add_argument("--direction", default="", help="题材方向,如 '三农园艺'")
    sp.add_argument("--profile", default="", help="使用的题材包路径,如 profiles/sannong.md")
    sp.add_argument("--topic", default="", help="选定的选题(一句话)")
    sp.add_argument("--base-dir", default="./output", help="项目根目录(默认 ./output)")
    sp.set_defaults(func=cmd_init_project)

    sp = sub.add_parser("gen-base-image",
                        help="文生图 - 基准图(默认 1080x1920 真 9:16)")
    sp.add_argument("--prompt-file", help="prompt 文本文件路径")
    sp.add_argument("--prompt", help="prompt 内联文本(与 --prompt-file 二选一)")
    sp.add_argument("--output", required=True, help="输出图片路径")
    sp.add_argument("--size", default=DEFAULT_IMAGE_SIZE,
                    help=f"图片尺寸(默认 {DEFAULT_IMAGE_SIZE},真 9:16;"
                         f"注意 1024x1536 是 2:3 会被视频裁切)")
    sp.add_argument("--size-fallback", default=IMAGE_SIZE_FALLBACK,
                    help=f"尺寸被模型拒绝时的回退值(默认 {IMAGE_SIZE_FALLBACK},传空字符串禁用)")
    sp.add_argument("--model", default=MODEL_IMAGE,
                    help=f"图片模型 ID,默认 {MODEL_IMAGE}")
    sp.set_defaults(func=cmd_gen_base_image)

    sp = sub.add_parser("gen-variant-image",
                        help="图生图编辑 - 变体镜头(基于基准图,保证一致性)")
    sp.add_argument("--base-image", required=True, help="基准图路径")
    sp.add_argument("--edit-prompt-file", help="edit prompt 文本文件路径")
    sp.add_argument("--edit-prompt", help="edit prompt 内联文本")
    sp.add_argument("--output", required=True, help="输出图片路径")
    sp.add_argument("--size", default=DEFAULT_IMAGE_SIZE,
                    help=f"图片尺寸(默认 {DEFAULT_IMAGE_SIZE},与基准图保持一致)")
    sp.add_argument("--size-fallback", default=IMAGE_SIZE_FALLBACK,
                    help="尺寸被拒时的回退值")
    sp.add_argument("--model", default=MODEL_IMAGE,
                    help=f"图片模型 ID,默认 {MODEL_IMAGE}")
    sp.set_defaults(func=cmd_gen_variant_image)

    sp = sub.add_parser("gen-video-clip",
                        help="单镜图生视频(异步轮询,默认带配音)")
    sp.add_argument("--image", required=True, help="作为首帧的分镜静图")
    sp.add_argument("--motion-prompt-file", help="motion prompt 文本文件路径")
    sp.add_argument("--motion-prompt", help="motion prompt 内联文本")
    sp.add_argument("--output", required=True, help="输出 mp4 路径")
    _add_video_job_args(sp)
    sp.set_defaults(func=cmd_gen_video_clip)

    sp = sub.add_parser("gen-video-batch",
                        help="多镜并发图生视频(读 jobs JSON,前台运行,进度走 stderr)")
    sp.add_argument("--jobs-json", required=True,
                    help="jobs JSON 数组文件,每项字段同 gen-video-clip 参数(下划线风格):"
                         "image / motion_prompt_file / output / voiceover_text / label 等")
    sp.add_argument("--max-workers", type=int, default=3,
                    help="并发任务数(默认 3)")
    _add_video_job_args(sp, with_voiceover_text=False)
    sp.set_defaults(func=cmd_gen_video_batch)

    sp = sub.add_parser("extract-frames",
                        help="ffmpeg 抽帧(firstlast=承接自检 / fps=对标拆解)")
    sp.add_argument("--video", required=True, help="输入视频路径")
    sp.add_argument("--output-dir", required=True, help="帧输出目录")
    sp.add_argument("--mode", default="firstlast", choices=["firstlast", "fps"],
                    help="抽帧模式(默认 firstlast)")
    sp.add_argument("--fps", type=float, default=1.0,
                    help="fps 模式下的抽帧频率(默认 1 帧/秒)")
    sp.set_defaults(func=cmd_extract_frames)

    sp = sub.add_parser("concat-clips",
                        help="ffmpeg 把多段 mp4 无缝拼接成成片")
    sp.add_argument("--clips", nargs="+", required=True,
                    help="按播放顺序排列的 mp4 文件路径(空格分隔,2 段以上)")
    sp.add_argument("--output", required=True, help="输出拼接后的 mp4 路径")
    sp.set_defaults(func=cmd_concat_clips)

    sp = sub.add_parser("build-storyboard",
                        help="生成最终交付 storyboard.xlsx")
    sp.add_argument("--shots-json", required=True,
                    help="分镜数据 JSON 文件,数组结构,每项含 voiceover/subtitle/prev_state/...")
    sp.add_argument("--output", required=True, help="输出 xlsx 路径")
    sp.set_defaults(func=cmd_build_storyboard)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
